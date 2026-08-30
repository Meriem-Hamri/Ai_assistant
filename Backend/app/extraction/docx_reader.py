from docx import Document as DocxDocument
from pathlib import Path

from app.models.document import Document, DocumentPage


def extract_text_from_docx(file_path: str) -> Document:
    """
    Extrait le texte d'un document Word (.docx).

    Raises:
        FileNotFoundError: Si le fichier n'existe pas.
        ValueError: Si le fichier ne peut pas être lu.
    """

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Le fichier '{file_path}' est introuvable."
        )

    try:
        docx = DocxDocument(path)

        text = "\n".join(
            paragraph.text
            for paragraph in docx.paragraphs
            if paragraph.text.strip()
        )

        document = Document(
            filename=path.name
        )

        document.pages.append(
            DocumentPage(
                page_number=1,
                text=text
            )
        )
        document.metadata["type"] = "docx"

        return document

    except Exception as e:
        raise ValueError(
            f"Impossible de lire le document Word : {e}"
        )